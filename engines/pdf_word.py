"""Local layout-aware PDF-to-DOCX engine, launched and supervised by Rust."""
import argparse
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault('OMP_NUM_THREADS', '2')
os.environ.setdefault('OPENBLAS_NUM_THREADS', '2')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('input')
    parser.add_argument('output')
    parser.add_argument('--quality', type=int, default=100)
    args = parser.parse_args()
    if Path(args.input).suffix.lower() != '.pdf' or Path(args.output).suffix.lower() in ['.png', '.jpg', '.zip', '.xlsx', '.txt']:
        from pdf_exports import convert
        convert(args.input, args.output, args.quality)
        return
    if Path(args.output).suffix.lower() in ['.pptx', '.svg']:
        from pdf_slides import convert
        convert(args.input, args.output)
        return
    import cv2
    import numpy as np
    import pymupdf
    from pdf2docx import Converter
    from pdf2docx.page.RawPageFitz import RawPageFitz

    cv2.setNumThreads(2)
    root = Path(getattr(sys, '_MEIPASS', Path(__file__).parent))
    data = root / 'tessdata'
    if not data.is_dir():
        data = Path(__file__).parent / 'tessdata'
    print('dropform_progress=0.02', flush=True)

    with tempfile.TemporaryDirectory(prefix='dropform-word-') as temp:
        source = pymupdf.open(args.input)
        if source.needs_pass:
            raise ValueError('Unlock this PDF before conversion.')
        prepared = pymupdf.open()
        ocr_pages = set()
        for index, page in enumerate(source):
            text = page.get_text('text').strip()
            if len(text) < 8 and page.get_images() and data.is_dir():
                scale = min(200 / 72, 3500 / max(page.rect.width, page.rect.height, 1))
                pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), colorspace=pymupdf.csRGB, alpha=False)
                pixels = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3)
                gray = cv2.cvtColor(pixels, cv2.COLOR_RGB2GRAY)
                # A scan normally has a dominant paper/background tone. Do not OCR arbitrary photos.
                paper = max(float(np.mean(gray > 210)), float(np.mean(gray < 45)))
                if paper < 0.45:
                    prepared.insert_pdf(source, from_page=index, to_page=index)
                    print(f'dropform_progress={0.05 + 0.35 * (index+1)/max(len(source),1):.4f}', flush=True)
                    continue
                scan = pymupdf.open('pdf', pix.pdfocr_tobytes(language='eng+ind', tessdata=str(data)))
                recognized = scan[0].get_text('text').strip()
                if len(recognized) >= 8:
                    prepared.insert_pdf(scan)
                    ocr_pages.add(index)
                else:
                    prepared.insert_pdf(source, from_page=index, to_page=index)
                scan.close()
            else:
                prepared.insert_pdf(source, from_page=index, to_page=index)
            print(f'dropform_progress={0.05 + 0.35 * (index+1)/max(len(source),1):.4f}', flush=True)
        source.close()
        pdf = Path(temp) / 'prepared.pdf'
        prepared.save(pdf, garbage=3, deflate=True)
        prepared.close()

        # OCR pages use the recognized text layer, rather than exporting the scan as a picture.
        original_text = RawPageFitz._preprocess_text
        original_images = RawPageFitz._preprocess_images

        def read_text(self, **settings):
            settings['ocr'] = 2 if self.page_engine.number in ocr_pages else 0
            return original_text(self, **settings)

        def read_images(self, **settings):
            settings['ocr'] = 2 if self.page_engine.number in ocr_pages else 0
            return original_images(self, **settings)

        RawPageFitz._preprocess_text = read_text
        RawPageFitz._preprocess_images = read_images
        converter = Converter(str(pdf))
        try:
            converter.convert(args.output, multi_processing=False, ignore_page_error=False,
                              line_break_width_ratio=1.0)
        finally:
            converter.close()
        from docx import Document
        from docx.oxml.ns import qn
        result = Document(args.output)
        for section in result.sections:
            columns = section._sectPr.find(qn('w:cols'))
            if columns is not None:
                columns.set(qn('w:space'), '0')
        from docx.oxml import OxmlElement
        from docx.enum.section import WD_SECTION_START
        sections=list(result.sections)
        for index, section in enumerate(sections):
            if index and section.start_type == WD_SECTION_START.NEW_COLUMN:
                previous=sections[index-1]._sectPr
                owner=previous.getparent()
                if owner.tag == qn('w:pPr'):
                    paragraph=owner.getparent()
                    start=previous.find(qn('w:type'))
                    section.start_type=WD_SECTION_START.CONTINUOUS
                    if start is not None:
                        section._sectPr.get_or_add_type().set(qn('w:val'),start.get(qn('w:val')))
                    owner.remove(previous)
                    run=OxmlElement('w:r');br=OxmlElement('w:br');br.set(qn('w:type'),'column');run.append(br);paragraph.append(run)
                    spacing = owner.find(qn('w:spacing'))
                    if spacing is None:
                        spacing = OxmlElement('w:spacing'); owner.append(spacing)
                    for key, value in [('before', '0'), ('after', '0'), ('line', '1'), ('lineRule', 'exact')]:
                        spacing.set(qn('w:'+key), value)
        fonts = {'Helvetica': 'Arial', 'Times-Roman': 'Times New Roman', 'Courier': 'Courier New'}
        for node in result.element.iter(qn('w:rFonts')):
            for key in ['ascii', 'hAnsi', 'eastAsia']:
                name = node.get(qn('w:'+key))
                if name in fonts:
                    node.set(qn('w:'+key), fonts[name])
        from docx.shared import Pt
        for paragraph in result.paragraphs:
            indent = paragraph.paragraph_format.right_indent
            if indent and indent > 0:
                # Leave room for PDF versus Word font-metric rounding at the right edge.
                paragraph.paragraph_format.right_indent = max(0, indent - Pt(6))
        def normalize_tables(tables):
            for table in tables:
                if table.rows and len(table.rows[0].cells) == len(table.columns):
                    for column, cell in zip(table.columns, table.rows[0].cells):
                        if cell.width:
                            column.width = cell.width
                for row in table.rows:
                    for cell in row.cells:
                        normalize_tables(cell.tables)
        normalize_tables(result.tables)
        result.save(args.output)
        print('dropform_progress=1.0', flush=True)


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(f'Conversion failed: {error}', file=sys.stderr, flush=True)
        sys.exit(1)
