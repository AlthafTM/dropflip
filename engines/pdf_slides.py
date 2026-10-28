"""PDF pages to positioned PowerPoint text over preserved page artwork."""
from io import BytesIO
from pathlib import Path
import math
import pymupdf


def convert(source, output):
    document = pymupdf.open(source)
    if document.needs_pass:
        raise ValueError('Unlock this PDF before conversion.')
    if not len(document):
        raise ValueError('This PDF has no pages.')
    suffix = Path(output).suffix.lower()
    if suffix == '.svg':
        Path(output).write_text(document[0].get_svg_image(text_as_path=True), encoding='utf-8')
        document.close()
        return
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    deck = Presentation()
    width, height = document[0].rect.width, document[0].rect.height
    deck.slide_width, deck.slide_height = Pt(width), Pt(height)
    for index, original in enumerate(document):
        # Normalize rotated pages before extracting coordinates and rendering artwork.
        page_copy = pymupdf.open()
        page_copy.insert_pdf(document, from_page=index, to_page=index)
        page = page_copy[0]
        page.remove_rotation()
        blocks = page.get_text('dict', flags=pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_IMAGES)['blocks']
        lines = [line for block in blocks if 'lines' in block for line in block['lines']]
        spans = [(line, span) for line in lines for span in line['spans'] if span['text'].strip()]
        # Hidden OCR text stays hidden; a scan remains a faithful page picture.
        traces = page.get_texttrace()
        hidden = [(pymupdf.Rect(trace['bbox']),trace['chars'][0][2][1]) for trace in traces if trace['type']==3 and trace['chars']]
        if traces and all(trace['type']==3 for trace in traces):
            spans = []
        elif hidden:
            spans = [(line,span) for line,span in spans if sum((box & pymupdf.Rect(span['bbox'])).width for box,baseline in hidden if abs(baseline-span['origin'][1])<1) < (span['bbox'][2]-span['bbox'][0])*.45]
        notes = page.get_text('text', sort=True)
        if any(abs(line.get('dir', (1, 0))[1])>.01 or line.get('dir', (1, 0))[0]<0 for line, _ in spans):
            spans = []  # Preserve rotated typography as artwork instead of placing incorrect text boxes.
        for line, span in spans:
            page.add_redact_annot(span['bbox'], fill=False, cross_out=False)
        if spans:
            page.apply_redactions(images=0, graphics=0, text=0)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = notes
        scale = min(width / page.rect.width, height / page.rect.height)
        left, top = (width-page.rect.width*scale)/2, (height-page.rect.height*scale)/2
        raster_scale = min(2.0, 2600/max(page.rect.width, page.rect.height))
        pix = page.get_pixmap(matrix=pymupdf.Matrix(raster_scale, raster_scale), alpha=False)
        slide.shapes.add_picture(BytesIO(pix.tobytes('png')), Pt(left), Pt(top), Pt(page.rect.width*scale), Pt(page.rect.height*scale))
        for line, span in spans:
            size = span['size'] * scale
            x, baseline = span['origin']
            box = slide.shapes.add_textbox(Pt(left+x*scale), Pt(top+baseline*scale-size*.86),
                                          Pt(max(2, (span['bbox'][2]-span['bbox'][0])*scale+size*.6)), Pt(size*1.6))
            frame = box.text_frame
            frame.margin_left = frame.margin_top = frame.margin_right = frame.margin_bottom = 0
            frame.word_wrap = False
            frame.auto_size = MSO_AUTO_SIZE.NONE
            frame.vertical_anchor = MSO_ANCHOR.TOP
            paragraph = frame.paragraphs[0]
            paragraph.space_before = paragraph.space_after = 0
            paragraph.line_spacing = 1.0
            run = paragraph.add_run()
            run.text = span['text']
            font = run.font
            name = span['font'].split('+')[-1]
            font.name = 'Arial' if 'Helvetica' in name or 'Arial' in name else 'Times New Roman' if 'Times' in name else 'Courier New' if 'Courier' in name else name
            font.size = Pt(size)
            font.bold, font.italic = bool(span['flags'] & 16), bool(span['flags'] & 2)
            color = span['color']
            font.color.rgb = RGBColor((color>>16)&255, (color>>8)&255, color&255)
            direction = line.get('dir', (1, 0))
            if abs(direction[1]) > .01 or direction[0] < 0:
                box.rotation = math.degrees(math.atan2(direction[1], direction[0])) % 360
        page_copy.close()
        print(f'dropform_progress={(index+1)/len(document):.4f}', flush=True)
    document.close()
    deck.save(output)
