"""Exercise resource boundaries through the real CLI using synthetic files."""
import json
import subprocess
import sys
import tempfile
import threading
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import pymupdf
from docx import Document
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
EXE = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / 'target/release/dropflip.exe'
checks = 0


def conversion(source, extension, destination, succeeds=True):
    global checks
    process = subprocess.run([str(EXE), '--convert', str(source), extension, str(destination)],
                             capture_output=True, text=True, encoding='utf-8', timeout=90)
    if succeeds:
        assert process.returncode == 0, process.stderr
        result = Path(json.loads(process.stdout)['path'])
        assert result.is_file() and result.stat().st_size
    else:
        assert process.returncode != 0, f'Unexpected success: {source.name}'
        assert not list(destination.iterdir()), f'Failed output published: {destination}'
        result = process.stderr
    checks += 1
    return result


with tempfile.TemporaryDirectory(prefix='dropflip-security-') as directory:
    work = Path(directory)
    source = work / 'inputs'
    source.mkdir()
    image = work / 'private.png'
    Image.new('RGB', (64, 48), 'red').save(image)
    ordinary = source / 'ordinary.md'
    ordinary.write_text('# Conversion control\n\nKeep **bold text** and this paragraph.\n', encoding='utf-8')
    for extension in ['html', 'docx', 'pdf', 'md.zip']:
        conversion(ordinary, extension, work / extension)
    word = source / 'embedded.docx'
    document = Document()
    document.add_paragraph('Embedded image control')
    document.add_picture(str(image))
    document.save(word)
    html = conversion(word, 'html', work / 'embedded-html').read_text(encoding='utf-8')
    assert 'data:image/' in html and 'Embedded image control' in html
    with zipfile.ZipFile(conversion(word, 'md.zip', work / 'embedded-bundle')) as archive:
        assert any(name.startswith('media/') for name in archive.namelist())
    local = source / 'outside.md'
    local.write_text('# Outside resource\n\n![outside](../private.png)\n', encoding='utf-8')
    conversion(local, 'html', work / 'denied-local', False)
    inline = source / 'include.html'
    inline.write_text(f'<html><head><title>Test</title></head><body><iframe src="{image.as_uri()}"></iframe></body></html>', encoding='utf-8')
    conversion(inline, 'docx', work / 'denied-iframe', False)
    requests = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            requests.append(self.path)
            self.send_response(200)
            self.end_headers()
            self.wfile.write(image.read_bytes())

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        remote = source / 'remote.md'
        remote.write_text(f'# Network resource\n\n![remote](http://127.0.0.1:{server.server_port}/pixel.png)\n', encoding='utf-8')
        conversion(remote, 'html', work / 'denied-network', False)
        assert requests == [], requests
    finally:
        server.shutdown()
        server.server_close()
        thread.join()
    raw = source / 'raw.md'
    raw.write_text('# Typst boundary\n\n```{=typst}\n#read("../private.png")\n```\n', encoding='utf-8')
    conversion(raw, 'pdf', work / 'denied-typst', False)
    huge = source / 'oversized.pdf'
    with pymupdf.open() as document:
        page = document.new_page(width=40000, height=40000)
        page.insert_image(pymupdf.Rect(0, 0, 100, 100), filename=str(image))
        document.save(huge)
    error = conversion(huge, 'png', work / 'denied-raster', False)
    assert 'too large to render safely' in error, error
    error = conversion(huge, 'txt', work / 'denied-ocr', False)
    assert 'too large to render safely' in error, error
    svg = source / 'oversized.svg'
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100000" height="100000"><rect width="100000" height="100000" fill="red"/></svg>', encoding='utf-8')
    error = conversion(svg, 'png', work / 'denied-svg', False)
    assert 'too large to render safely' in error, error

print(f'{checks} real-output security checks passed')
